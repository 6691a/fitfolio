import asyncio
import logging
import re
import uuid
from pathlib import Path
from typing import cast

import aiofiles
import fitz
import trafilatura
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from fastapi import UploadFile

from app.ai.classification.document import DocumentClassificationError, DocumentClassifierProtocol
from app.ai.extraction import (
    StructuredExtractionError,
    extract_job_posting_structured,
    extract_resume_structured,
)
from app.ai.vision import extract_image_content
from app.config.settings import settings
from app.crawlers.job_postings import JobPostingCrawler
from app.repositories.documents import DocumentsRepository
from app.repositories.profiles import ProfilesRepository
from app.schemas.documents import (
    SUPPORTED_MIME_TYPES,
    DocumentFormat,
    DocumentInput,
    DocumentKind,
    FileInput,
    JobPostingExtractDebug,
    ParseApplicationAccepted,
    ParsedDocument,
    ParseJobStatus,
    ParseStatus,
    ResumeExtractDebug,
    TextInput,
    UrlInput,
)
from app.schemas.profiles import JobPostingProfileData, ResumeProfileData
from app.services.profiles import build_profile_data
from app.utils import clean_text

_UPLOAD_SUBDIRS: dict[DocumentKind, str] = {
    DocumentKind.RESUME: "resumes",
    DocumentKind.JOB_POSTING: "job_postings",
}

logger = logging.getLogger(__name__)

# 디버그: 비전에 보낼 이미지를 디스크에 저장.
_DEBUG_IMAGE_DIR = settings.UPLOAD_DIR / settings.DOCUMENT_DEBUG_IMAGE_SUBDIR


class DocumentFileParseError(Exception):
    pass


class UnsupportedDocumentFormatError(Exception):
    pass


class InsufficientJobContentError(Exception):
    pass


class FileTooLargeError(Exception):
    pass


class DocumentService:
    _CHUNK_SIZE = 1024 * 1024

    @staticmethod
    def _save_debug_image(data: bytes, ext: str, tag: str) -> None:
        """디버그가 켜져 있으면 이미지를 디버그 디렉터리에 저장한다.

        Args:
            data: 저장할 이미지 바이트.
            ext: 파일 확장자(예: png, jpg).
            tag: 파일명 접두 태그.
        """
        if not settings.DOCUMENT_DEBUG_ENABLED:
            return
        path = _DEBUG_IMAGE_DIR / f"{tag}-{uuid.uuid4().hex[:8]}.{ext}"
        path.write_bytes(data)
        logger.info("디버그 이미지 저장: %s", path)

    @staticmethod
    def _save_debug_text(text: str, tag: str, ext: str = "html") -> None:
        """디버그가 켜져 있으면 텍스트를 디버그 디렉터리에 저장한다.

        Args:
            text: 저장할 텍스트.
            tag: 파일명 접두 태그.
            ext: 파일 확장자(기본 html).
        """
        if not settings.DOCUMENT_DEBUG_ENABLED:
            return
        path = _DEBUG_IMAGE_DIR / f"{tag}-{uuid.uuid4().hex[:8]}.{ext}"
        path.write_text(text, encoding="utf-8")
        logger.info("디버그 HTML 저장: %s (len=%d)", path, len(text))

    def __init__(
        self,
        documents_repository: DocumentsRepository,
        job_posting_crawler: JobPostingCrawler | None = None,
        profiles_repository: ProfilesRepository | None = None,
        document_classifier: DocumentClassifierProtocol | None = None,
    ) -> None:
        """협력자(레포지토리·크롤러·분류기)를 주입받아 보관한다.

        Args:
            documents_repository: 문서 영속화 레포지토리.
            job_posting_crawler: 채용공고 URL 크롤러(URL 파싱에 필요).
            profiles_repository: 프로필 영속화 레포지토리(프로필 저장에 필요).
            document_classifier: 문서 유형 판별기(유형 검증에 필요).
        """
        self._documents_repository = documents_repository
        self._job_posting_crawler = job_posting_crawler
        self._profiles_repository = profiles_repository
        self._document_classifier = document_classifier

    async def process(self, document_id: int) -> None:
        """문서를 파싱→유형 검증→완료 처리/프로필 저장까지 오케스트레이션한다.

        파싱 실패나 유형 불일치 시 문서를 실패로 표시하고 중단한다.

        Args:
            document_id: 처리할 문서 ID(없으면 아무 것도 하지 않음).
        """
        document = await self._documents_repository.get(document_id)
        if document is None:
            return

        await self._documents_repository.mark_started(document_id)

        document_input = self._build_document_input(document)

        try:
            parsed = await self.parse(document_input)
        except Exception as exc:
            logger.exception("문서 파싱 실패 document_id=%s", document_id)
            await self._documents_repository.mark_failed(document_id, error=str(exc))
            return

        if not await self.verify_kind(
            document_id=document_id,
            expected_kind=DocumentKind(document.document_type),
            parsed=parsed,
        ):
            return

        await self._documents_repository.mark_done(
            document_id,
            extracted_text=parsed.extracted_text,
            metadata=parsed.metadata,
        )
        await self.store_profile(document, parsed)

    def _build_document_input(self, document) -> DocumentInput:
        """DB 문서 레코드를 형식에 맞는 파싱 입력 스키마로 변환한다.

        Args:
            document: DB에서 조회한 문서 레코드.

        Returns:
            형식에 맞는 UrlInput/TextInput/FileInput.

        Raises:
            ValueError: 지원하지 않는 문서 형식일 때.
        """
        document_format = DocumentFormat(document.format)
        match document_format:
            case DocumentFormat.URL:
                return UrlInput(
                    document_type=document.document_type,
                    input_type=DocumentFormat.URL,
                    url=document.source_url,
                )
            case DocumentFormat.TEXT:
                return TextInput(
                    document_type=document.document_type,
                    input_type=DocumentFormat.TEXT,
                    text=document.extracted_text or "",
                )
            case DocumentFormat.PDF | DocumentFormat.DOCX | DocumentFormat.IMAGE | DocumentFormat.PPT:
                return FileInput(
                    document_type=document.document_type,
                    input_type=document_format,
                    file_path=document.file_path,
                    file_name=document.file_name,
                    content_type=document.content_type,
                )

        raise ValueError(f"Unsupported document format: {document_format}")

    async def verify_kind(
        self,
        *,
        document_id: int,
        expected_kind: DocumentKind,
        parsed: ParsedDocument,
    ) -> bool:
        """파싱 결과가 기대 문서 종류와 맞는지 판별하고 불일치 시 실패 처리한다.

        판별 결과는 parsed.metadata['document_classification']에 기록된다.

        Args:
            document_id: 검증 대상 문서 ID.
            expected_kind: 기대하는 문서 종류.
            parsed: 검증할 파싱 결과(메타데이터가 갱신됨).

        Returns:
            기대 종류와 일치하면 True, 아니면(또는 판별 실패) False.

        Raises:
            RuntimeError: 분류기가 주입되지 않은 경우.
        """
        if self._document_classifier is None:
            raise RuntimeError("DocumentClassifier must be injected through Container.document_classifier")

        try:
            classification = await self._document_classifier.classify(parsed.extracted_text, expected_kind)
        except DocumentClassificationError as exc:
            logger.exception("문서 유형 판단 실패 document_id=%s", document_id)
            await self._documents_repository.mark_failed(document_id, error=f"문서 유형 판단에 실패했습니다: {exc}")
            return False

        parsed.metadata["document_classification"] = classification.model_dump(mode="json", exclude_none=True)
        if classification.is_expected:
            return True

        await self._documents_repository.mark_failed(
            document_id,
            error=classification.failure_reason or "업로드한 문서가 요청한 문서 종류와 일치하지 않습니다.",
        )
        return False

    async def store_profile(self, document, parsed: ParsedDocument) -> None:
        """파싱 결과로 이력서/채용공고 프로필을 생성해 저장한다.

        Args:
            document: 프로필이 속한 문서 레코드.
            parsed: 프로필을 만들 파싱 결과.

        Raises:
            RuntimeError: 프로필 레포지토리가 주입되지 않은 경우.
        """
        if self._profiles_repository is None:
            raise RuntimeError("ProfilesRepository must be injected through Container.profiles_repository")

        profile = build_profile_data(DocumentKind(document.document_type), parsed)
        if isinstance(profile, ResumeProfileData):
            await self._profiles_repository.upsert_resume_profile(document_id=document.id, profile=profile)
            return

        if isinstance(profile, JobPostingProfileData):
            await self._profiles_repository.upsert_job_posting_profile(document_id=document.id, profile=profile)

    async def get_parse_status(self, document_id: int) -> ParseJobStatus | None:
        """문서 파싱 진행 상태를 조회하고 완료 시 결과까지 만들어 반환한다.

        Args:
            document_id: 상태를 조회할 문서 ID.

        Returns:
            상태(및 DONE이면 결과, FAILED면 오류)를 담은 ParseJobStatus. 문서가 없으면 None.

        Raises:
            ValueError: 레코드의 형식이 지원하지 않는 값일 때.
        """
        record = await self._documents_repository.get(document_id)
        if record is None:
            return None

        if record.status == ParseStatus.DONE:
            document_format = DocumentFormat(record.format)
            match document_format:
                case DocumentFormat.URL:
                    original_input = UrlInput(
                        document_type=record.document_type,
                        input_type=DocumentFormat.URL,
                        url=record.source_url,
                    )
                case DocumentFormat.PDF | DocumentFormat.DOCX | DocumentFormat.IMAGE | DocumentFormat.PPT:
                    assert record.file_path is not None
                    assert record.file_name is not None
                    assert record.content_type is not None
                    original_input = FileInput(
                        document_type=record.document_type,
                        input_type=document_format,
                        file_path=record.file_path,
                        file_name=record.file_name,
                        content_type=record.content_type,
                    )
                case _:
                    raise ValueError(f"Unsupported document format: {document_format}")

            result = ParsedDocument(
                original_input=original_input,
                extracted_text=record.extracted_text or "",
                metadata=record.extracted_metadata or {},
            )
            return ParseJobStatus(document_id=document_id, status=ParseStatus.DONE, result=result)

        if record.status == ParseStatus.FAILED:
            return ParseJobStatus(document_id=document_id, status=ParseStatus.FAILED, error=record.error)

        return ParseJobStatus(document_id=document_id, status=record.status)

    async def save(
        self,
        format: DocumentFormat,
        file: UploadFile,
        *,
        document_type: DocumentKind,
        suffix: str,
        max_bytes: int,
    ) -> Path:
        """업로드 파일을 검증하고 UUID 이름으로 디스크에 저장한다.

        Args:
            format: 파일 입력 형식(PDF/DOCX 등).
            file: 업로드된 파일 객체.
            document_type: 저장 경로를 결정할 문서 종류.
            suffix: 저장 파일 확장자(예: .pdf).
            max_bytes: 허용 최대 바이트 수.

        Returns:
            저장된 파일의 경로.

        Raises:
            UnsupportedDocumentFormatError: 지원하지 않는 파일 형식일 때.
            FileTooLargeError: 허용 크기를 초과할 때.
        """
        if file.content_type not in SUPPORTED_MIME_TYPES.get(format, frozenset()):
            raise UnsupportedDocumentFormatError("지원하지 않는 파일 형식입니다")

        stored_dir = settings.UPLOAD_DIR / _UPLOAD_SUBDIRS[document_type]
        stored_dir.mkdir(parents=True, exist_ok=True)
        stored_path = stored_dir / f"{uuid.uuid4()}{suffix}"
        size = 0

        try:
            async with aiofiles.open(stored_path, "wb") as out_file:
                while chunk := await file.read(self._CHUNK_SIZE):
                    size += len(chunk)
                    if size > max_bytes:
                        raise FileTooLargeError("파일이 너무 큽니다")
                    await out_file.write(chunk)
        except FileTooLargeError:
            stored_path.unlink(missing_ok=True)
            raise

        return stored_path

    async def request_parse(
        self,
        *,
        document_type: DocumentKind,
        format: DocumentFormat,
        file_name: str,
        file_path: str,
        content_type: str,
    ) -> int:
        """파일 문서 레코드를 만들고 비동기 파싱 작업을 큐에 등록한다.

        Args:
            document_type: 문서 종류(이력서/채용공고).
            format: 파일 입력 형식.
            file_name: 업로드 원본 파일명.
            file_path: 저장된 파일 경로.
            content_type: 파일 MIME 타입.

        Returns:
            생성된 문서의 ID.
        """
        from app.tasks.documents import task_parse_document

        record = await self._documents_repository.create(
            document_type=document_type,
            format=format,
            file_name=file_name,
            file_path=file_path,
            content_type=content_type,
        )
        task_parse_document.delay(record.id)
        return record.id

    async def request_parse_url(self, *, document_type: DocumentKind, url: str) -> int:
        """URL 문서 레코드를 만들고 비동기 파싱 작업을 큐에 등록한다.

        Args:
            document_type: 문서 종류(보통 채용공고).
            url: 파싱할 채용공고 URL.

        Returns:
            생성된 문서의 ID.
        """
        from app.tasks.documents import task_parse_document

        record = await self._documents_repository.create(
            document_type=document_type,
            format=DocumentFormat.URL,
            source_url=url,
        )
        task_parse_document.delay(record.id)
        return record.id

    async def request_parse_text(self, *, document_type: DocumentKind, text: str) -> int:
        """텍스트 문서 레코드를 만들고 비동기 파싱 작업을 큐에 등록한다.

        Args:
            document_type: 문서 종류(이력서/채용공고).
            text: 파싱할 본문 텍스트.

        Returns:
            생성된 문서의 ID.
        """
        from app.tasks.documents import task_parse_document

        record = await self._documents_repository.create(
            document_type=document_type,
            format=DocumentFormat.TEXT,
            extracted_text=text,
        )
        task_parse_document.delay(record.id)
        return record.id

    async def request_application_parse(
        self,
        *,
        resume_format: DocumentFormat,
        resume_file: UploadFile,
        job_posting_format: DocumentFormat,
        job_posting_file: UploadFile | None = None,
        job_posting_url: str | None = None,
        job_posting_text: str | None = None,
    ) -> ParseApplicationAccepted:
        """이력서와 채용공고를 저장·등록하고 두 파싱 작업을 동시에 큐에 올린다.

        입력은 컨트롤러에서 형식/필수값 검증을 마친 값이라고 가정한다.

        Args:
            resume_format: 이력서 입력 형식.
            resume_file: 업로드된 이력서 파일.
            job_posting_format: 채용공고 입력 형식.
            job_posting_file: 업로드된 채용공고 파일(파일 입력일 때).
            job_posting_url: 채용공고 URL(URL 입력일 때).
            job_posting_text: 채용공고 본문(텍스트 입력일 때).

        Returns:
            등록된 이력서/채용공고 문서 ID를 담은 ParseApplicationAccepted.
        """
        resume_document_id, job_posting_document_id = await asyncio.gather(
            self._request_resume_parse(resume_format, resume_file),
            self._request_job_posting_parse(job_posting_format, job_posting_file, job_posting_url, job_posting_text),
        )
        return ParseApplicationAccepted(
            resume_document_id=resume_document_id,
            job_posting_document_id=job_posting_document_id,
        )

    async def _request_resume_parse(self, resume_format: DocumentFormat, resume_file: UploadFile) -> int:
        """이력서 파일을 저장하고 파싱 작업을 등록한다.

        Args:
            resume_format: 이력서 입력 형식.
            resume_file: 업로드된 이력서 파일(파일명·content_type 검증 완료).

        Returns:
            생성된 이력서 문서의 ID.
        """
        assert resume_file.filename is not None
        assert resume_file.content_type is not None
        resume_path = await self.save(
            resume_format,
            resume_file,
            document_type=DocumentKind.RESUME,
            suffix=f".{resume_format.value}",
            max_bytes=settings.MAX_PDF_BYTES,
        )
        return await self.request_parse(
            document_type=DocumentKind.RESUME,
            format=resume_format,
            file_name=resume_file.filename,
            file_path=str(resume_path),
            content_type=resume_file.content_type,
        )

    async def _request_job_posting_parse(
        self,
        job_posting_format: DocumentFormat,
        job_posting_file: UploadFile | None,
        job_posting_url: str | None,
        job_posting_text: str | None,
    ) -> int:
        """채용공고를 형식(URL/텍스트/파일)에 맞게 저장·등록한다.

        Args:
            job_posting_format: 채용공고 입력 형식.
            job_posting_file: 업로드된 채용공고 파일(파일 입력일 때).
            job_posting_url: 채용공고 URL(URL 입력일 때).
            job_posting_text: 채용공고 본문(텍스트 입력일 때).

        Returns:
            생성된 채용공고 문서의 ID.
        """
        if job_posting_format == DocumentFormat.URL:
            assert job_posting_url is not None
            return await self.request_parse_url(document_type=DocumentKind.JOB_POSTING, url=job_posting_url)
        if job_posting_format == DocumentFormat.TEXT:
            assert job_posting_text is not None
            return await self.request_parse_text(document_type=DocumentKind.JOB_POSTING, text=job_posting_text)

        assert job_posting_file is not None
        assert job_posting_file.filename is not None
        assert job_posting_file.content_type is not None
        job_posting_path = await self.save(
            job_posting_format,
            job_posting_file,
            document_type=DocumentKind.JOB_POSTING,
            suffix=f".{job_posting_format.value}",
            max_bytes=settings.MAX_PDF_BYTES,
        )
        return await self.request_parse(
            document_type=DocumentKind.JOB_POSTING,
            format=job_posting_format,
            file_name=job_posting_file.filename,
            file_path=str(job_posting_path),
            content_type=job_posting_file.content_type,
        )

    async def parse(self, input: DocumentInput) -> ParsedDocument:
        """입력 형식에 맞는 파서로 분기해 문서를 파싱한다.

        Args:
            input: 파싱할 문서 입력(형식별 스키마).

        Returns:
            추출 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            UnsupportedDocumentFormatError: 아직 지원하지 않는 형식(IMAGE/PPT)일 때.
            ValueError: 알 수 없는 입력 형식일 때.
        """
        match input.input_type:
            case DocumentFormat.PDF:
                return await self.pdf_to_text(input)
            case DocumentFormat.URL:
                return await self.url_to_text(input)
            case DocumentFormat.TEXT:
                return await self.text_to_parsed(input)
            case DocumentFormat.DOCX:
                return await self.docx_to_text(input)
            case DocumentFormat.IMAGE | DocumentFormat.PPT:
                raise UnsupportedDocumentFormatError(
                    "아직 지원하지 않는 파일 형식입니다. 현재는 PDF/DOCX 이력서와 URL/텍스트 채용공고를 사용해주세요."
                )

        raise ValueError(f"Unsupported input type: {input.input_type}")

    async def pdf_to_text(self, input: FileInput) -> ParsedDocument:
        """PDF의 본문 텍스트를 추출하고 임베드 이미지는 비전으로 보조 추출한다.

        Args:
            input: 파싱할 PDF 파일 정보.

        Returns:
            본문+이미지 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            DocumentFileParseError: PDF 파일을 열거나 읽지 못한 경우.
        """
        image_texts: list[str] = []
        try:
            with fitz.open(input.file_path) as doc:
                pages = [cast(str, page.get_text("text")) for page in doc]
                for page in doc:
                    for img in page.get_images(full=True):
                        base = doc.extract_image(img[0])
                        if base["width"] * base["height"] < settings.DOCUMENT_MIN_IMAGE_PIXELS:
                            logger.info(
                                "이미지 OCR 스킵: source=pdf reason=too_small width=%s height=%s bytes=%s",
                                base["width"],
                                base["height"],
                                len(base["image"]),
                            )
                            continue
                        self._save_debug_image(base["image"], base["ext"], "pdf")
                        mime = "image/jpeg" if base["ext"] in {"jpg", "jpeg"} else f"image/{base['ext']}"
                        result = await extract_image_content(base["image"], mime, input.document_type)
                        if result.relevant and result.content:
                            image_texts.append(result.content)
        except (fitz.FileDataError, fitz.EmptyFileError, FileNotFoundError, OSError, ValueError) as exc:
            raise DocumentFileParseError(
                "PDF 파일을 읽지 못했습니다. 실제 PDF 파일인지 확인한 뒤 다시 업로드해주세요."
            ) from exc

        text = "\n\n".join(pages)
        if image_texts:
            text += "\n\n" + "\n\n".join(image_texts)

        return ParsedDocument(
            original_input=input,
            extracted_text=text,
            metadata=await self._file_metadata(input, text, len(pages), image_texts),
        )

    async def docx_to_text(self, input: FileInput) -> ParsedDocument:
        """DOCX 파일을 텍스트로 추출해 ParsedDocument로 변환한다.

        Args:
            input: 파싱할 DOCX 파일 정보.

        Returns:
            추출 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            DocumentFileParseError: 실제 .docx가 아니거나 파일을 열 수 없을 때.
        """
        try:
            text = self._extract_docx_text(input.file_path)
        except PackageNotFoundError as exc:
            raise DocumentFileParseError(
                "DOCX 파일을 읽지 못했습니다. 실제 Word(.docx) 파일인지 확인한 뒤 다시 업로드해주세요."
            ) from exc

        return ParsedDocument(
            original_input=input,
            extracted_text=text,
            # docx는 페이지 개념이 없어 page_count=0, 이미지 OCR은 범위 외라 image_texts=[]
            metadata=await self._file_metadata(input, text, 0, []),
        )

    @staticmethod
    def _extract_docx_text(file_path: str) -> str:
        """DOCX의 본문 단락과 표 셀 텍스트를 합쳐 한 문자열로 만든다.

        Args:
            file_path: 읽을 .docx 파일 경로.

        Returns:
            단락과 표 텍스트를 줄바꿈으로 이어 붙인 문자열.
        """
        document = DocxDocument(file_path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        # ponytail: 본문 단락 뒤에 표 텍스트를 이어 붙임 — 단락/표 사이 원본 순서는 보존 안 함.
        # 순서가 중요해지면 docx body 요소(iter_block_items)로 교체.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append("\t".join(cells))
        return "\n".join(parts)

    async def url_to_text(self, input: UrlInput) -> ParsedDocument:
        """채용공고 URL을 크롤링해 본문 텍스트(또는 HTML)를 파싱한다.

        Args:
            input: 파싱할 채용공고 URL 입력.

        Returns:
            크롤링 결과를 정리한 ParsedDocument.

        Raises:
            RuntimeError: 채용공고 크롤러가 주입되지 않은 경우.
        """
        if self._job_posting_crawler is None:
            raise RuntimeError("JobPostingCrawler must be injected through Container.job_posting_crawler")

        result = await self._job_posting_crawler.crawl(str(input.url), input.document_type)
        if result.text is not None:
            return self.url_text_to_parsed(
                input,
                result.text,
                result.image_texts,
                result.structured_debug,
            )

        return await self.html_to_text(
            result.htmls,
            input,
            result.image_texts,
            result.structured_debug,
        )

    def url_text_to_parsed(
        self,
        input: UrlInput,
        text: str,
        image_texts: list[str] | None = None,
        structured_debug: JobPostingExtractDebug | None = None,
    ) -> ParsedDocument:
        """크롤러가 준 본문 텍스트를 정규화·검증해 ParsedDocument로 만든다.

        Args:
            input: 원본 채용공고 URL 입력.
            text: 크롤러가 추출한 본문 텍스트.
            image_texts: 이미지에서 추출한 보조 텍스트들.
            structured_debug: 구조화 추출 디버그 정보.

        Returns:
            정규화·절단된 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            InsufficientJobContentError: 추출 텍스트가 최소 길이에 못 미칠 때.
        """
        normalized = text.strip()
        image_texts = image_texts or []
        if image_texts:
            normalized = (normalized + "\n\n" + "\n\n".join(image_texts)).strip()

        if len(normalized) < settings.DOCUMENT_MIN_EXTRACTED_CHARS:
            raise InsufficientJobContentError(
                "이 URL에서 텍스트를 충분히 추출하지 못했습니다. 채용 공고 본문을 직접 복사해 붙여주세요."
            )

        truncated = normalized[: settings.DOCUMENT_MAX_EXTRACTED_CHARS]
        return ParsedDocument(
            original_input=input,
            extracted_text=truncated,
            metadata=self._url_metadata(
                input,
                truncated,
                image_texts,
                structured_debug,
            ),
        )

    def _url_metadata(
        self,
        input: UrlInput,
        text: str,
        image_texts: list[str],
        structured_debug: JobPostingExtractDebug | None = None,
    ) -> dict:
        """URL 채용공고용 메타데이터(원문 URL·글자 수·구조화 결과)를 만든다.

        Args:
            input: 원본 채용공고 URL 입력.
            text: 추출 본문 텍스트.
            image_texts: 이미지에서 추출한 보조 텍스트들.
            structured_debug: 구조화 추출 디버그 정보(있으면 포함).

        Returns:
            메타데이터 dict.
        """
        metadata: dict[str, object] = {
            "source_url": str(input.url),
            "char_count": len(text),
            "image_text_count": len(image_texts),
        }
        if structured_debug is not None:
            metadata["job_posting_extract"] = structured_debug.model_dump(
                mode="json",
                exclude_none=True,
            )
        return metadata

    async def _file_metadata(self, input: FileInput, text: str, page_count: int, image_texts: list[str]) -> dict:
        """파일 문서 메타데이터를 만들고 이력서면 구조화 추출을 포함한다.

        Args:
            input: 원본 파일 입력.
            text: 추출 본문 텍스트.
            page_count: 페이지 수(페이지 개념이 없으면 0).
            image_texts: 이미지에서 추출한 보조 텍스트들.

        Returns:
            메타데이터 dict(이력서면 resume_extract 포함).
        """
        metadata: dict[str, object] = {"page_count": page_count, "image_text_count": len(image_texts)}
        if input.document_type == DocumentKind.RESUME:
            payload = await self._resume_structured_payload(text, image_texts)
            metadata["resume_extract"] = self._dump_resume_debug_payload(payload)
        return metadata

    async def _text_metadata(self, input: TextInput, text: str) -> dict:
        """텍스트 입력 메타데이터를 만들고 종류별 구조화 추출을 포함한다.

        Args:
            input: 원본 텍스트 입력.
            text: 본문 텍스트.

        Returns:
            메타데이터 dict(이력서면 resume_extract, 채용공고면 job_posting_extract 포함).
        """
        metadata: dict[str, object] = {"char_count": len(text)}
        if input.document_type == DocumentKind.RESUME:
            payload = await self._resume_structured_payload(text)
            metadata["resume_extract"] = self._dump_resume_debug_payload(payload)
        elif input.document_type == DocumentKind.JOB_POSTING:
            payload = await self._job_posting_structured_payload(
                text,
                JobPostingExtractDebug(source="text", text=text),
            )
            metadata["job_posting_extract"] = payload.model_dump(mode="json", exclude_none=True)
        return metadata

    async def _resume_structured_payload(self, text: str, image_texts: list[str] | None = None) -> ResumeExtractDebug:
        """이력서 구조화 추출을 시도하고 실패 시 규칙 기반 폴백을 사용한다.

        Args:
            text: 이력서 본문 텍스트.
            image_texts: 이미지에서 추출한 보조 텍스트들.

        Returns:
            구조화된 ResumeExtractDebug.

        Raises:
            InsufficientJobContentError: 이력서/자기소개서/포트폴리오로 보이지 않을 때.
        """
        fallback = self._resume_debug_payload(text, image_texts)
        try:
            payload = await extract_resume_structured(text, fallback)
        except StructuredExtractionError:
            payload = fallback

        if image_texts and not payload.image_texts:
            payload.image_texts = image_texts
        if not payload.relevant:
            raise InsufficientJobContentError(
                payload.failure_reason or "이 문서는 이력서/자기소개서/포트폴리오로 보이지 않습니다."
            )
        return payload

    async def _job_posting_structured_payload(
        self, text: str, fallback: JobPostingExtractDebug
    ) -> JobPostingExtractDebug:
        """채용공고 구조화 추출을 시도하고 실패 시 폴백을 사용한다.

        Args:
            text: 채용공고 본문 텍스트.
            fallback: 추출 실패 시 사용할 기본 추출 결과.

        Returns:
            구조화된 JobPostingExtractDebug.

        Raises:
            InsufficientJobContentError: 채용공고로 보이지 않을 때.
        """
        try:
            payload = await extract_job_posting_structured(text, fallback)
        except StructuredExtractionError:
            payload = fallback

        if not payload.relevant:
            raise InsufficientJobContentError(payload.failure_reason or "이 문서는 채용공고로 보이지 않습니다.")
        return payload

    def _dump_resume_debug_payload(self, payload: ResumeExtractDebug) -> dict:
        """이력서 구조화 결과를 dict로 덤프하고 디버그 파일로도 저장한다.

        Args:
            payload: 덤프할 이력서 구조화 결과.

        Returns:
            None 필드를 제외한 JSON 직렬화 dict.
        """
        data = payload.model_dump(mode="json", exclude_none=True)
        self._save_debug_text(
            payload.model_dump_json(indent=2, exclude_none=True),
            "resume-extract",
            "json",
        )
        return data

    def _resume_debug_payload(self, text: str, image_texts: list[str] | None = None) -> ResumeExtractDebug:
        """규칙 기반(정규식/섹션 분할)으로 이력서 보수 추출 결과를 만든다.

        Args:
            text: 이력서 본문 텍스트.
            image_texts: 이미지에서 추출한 보조 텍스트들.

        Returns:
            규칙 기반으로 채운 ResumeExtractDebug(폴백용).
        """
        normalized = text.strip()
        sections = self._resume_sections(normalized)
        return ResumeExtractDebug(
            text=normalized,
            name=self._resume_labeled_value(normalized, ("성명", "이름", "Name")),
            email=self._resume_email(normalized),
            phone=self._resume_phone(normalized),
            self_introduction=self._resume_first_section(sections, ("자기소개", "소개", "Profile")),
            career_summary=self._resume_first_section(sections, ("경력요약", "요약", "Summary")),
            work_experiences=self._resume_section_list(sections, ("경력", "경력사항", "Experience")),
            projects=self._resume_section_list(sections, ("프로젝트", "Projects")),
            skills=self._resume_skills(sections),
            education=self._resume_section_list(sections, ("학력", "Education")),
            certifications=self._resume_section_list(sections, ("자격증", "Certifications", "Certificates")),
            image_texts=image_texts or [],
            raw_sections=sections,
        )

    def _resume_sections(self, text: str) -> dict[str, str]:
        """이력서 본문을 알려진 헤더 기준으로 섹션별 텍스트로 분할한다.

        Args:
            text: 이력서 본문 텍스트.

        Returns:
            헤더명 → 해당 섹션 내용 dict.
        """
        header_aliases = (
            "자기소개",
            "소개",
            "Profile",
            "경력요약",
            "요약",
            "Summary",
            "경력",
            "경력사항",
            "Experience",
            "프로젝트",
            "Projects",
            "기술스택",
            "보유기술",
            "Skills",
            "학력",
            "Education",
            "자격증",
            "Certifications",
            "Certificates",
        )
        header_pattern = "|".join(re.escape(header) for header in header_aliases)
        matches = list(
            re.finditer(
                rf"(?im)^\s*(?P<header>{header_pattern})\s*[:：]?\s*$",
                text,
            )
        )
        sections: dict[str, str] = {}
        for index, match in enumerate(matches):
            start = match.end()
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            content = text[start:end].strip()
            if content:
                sections[match.group("header")] = content
        return sections

    def _resume_labeled_value(self, text: str, labels: tuple[str, ...]) -> str | None:
        """'라벨: 값' 형식의 첫 줄에서 라벨에 해당하는 값을 찾는다.

        Args:
            text: 검색할 본문 텍스트.
            labels: 찾을 라벨 후보들(예: 성명/이름/Name).

        Returns:
            매칭된 값(공백 정리됨), 없으면 None.
        """
        label_pattern = "|".join(re.escape(label) for label in labels)
        match = re.search(rf"(?im)^\s*(?:{label_pattern})\s*[:：]\s*(.+?)\s*$", text)
        return clean_text(match.group(1)) if match else None

    def _resume_email(self, text: str) -> str | None:
        """본문에서 첫 이메일 주소를 추출한다.

        Args:
            text: 검색할 본문 텍스트.

        Returns:
            찾은 이메일 주소, 없으면 None.
        """
        match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
        return match.group(0) if match else None

    def _resume_phone(self, text: str) -> str | None:
        """본문에서 첫 한국 휴대폰 번호를 추출한다.

        Args:
            text: 검색할 본문 텍스트.

        Returns:
            찾은 전화번호, 없으면 None.
        """
        match = re.search(r"(?:\+82[-\s]?)?0?1[016789][-\s.]?\d{3,4}[-\s.]?\d{4}", text)
        return match.group(0) if match else None

    def _resume_first_section(self, sections: dict[str, str], labels: tuple[str, ...]) -> str | None:
        """섹션 dict에서 라벨 후보 중 처음 발견되는 섹션 내용을 반환한다.

        Args:
            sections: 헤더명 → 내용 섹션 dict.
            labels: 우선순위대로 찾을 라벨 후보들.

        Returns:
            첫 매칭 섹션 내용(공백 정리됨), 없으면 None.
        """
        for label in labels:
            value = sections.get(label)
            if value:
                return clean_text(value)
        return None

    def _resume_section_list(self, sections: dict[str, str], labels: tuple[str, ...]) -> list[str]:
        """섹션 내용을 줄/불릿 단위로 나눠 항목 리스트로 만든다.

        Args:
            sections: 헤더명 → 내용 섹션 dict.
            labels: 대상 섹션을 찾을 라벨 후보들.

        Returns:
            비어 있지 않은 항목 문자열 리스트.
        """
        value = self._resume_first_section(sections, labels)
        if not value:
            return []
        return [item for item in (clean_text(part) for part in re.split(r"\n+|[•·]", value)) if item]

    def _resume_skills(self, sections: dict[str, str]) -> list[str]:
        """기술스택 섹션을 구분자(쉼표/슬래시 등)로 나눠 스킬 리스트로 만든다.

        Args:
            sections: 헤더명 → 내용 섹션 dict.

        Returns:
            비어 있지 않은 스킬 문자열 리스트.
        """
        value = self._resume_first_section(sections, ("기술스택", "보유기술", "Skills"))
        if not value:
            return []
        return [skill for skill in (clean_text(part) for part in re.split(r"[,/|·•\n]+", value)) if skill]

    async def text_to_parsed(self, input: TextInput) -> ParsedDocument:
        """붙여넣은 텍스트를 검증·절단해 ParsedDocument로 만든다.

        Args:
            input: 파싱할 텍스트 입력.

        Returns:
            절단된 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            InsufficientJobContentError: 본문이 최소 길이에 못 미칠 때.
        """
        text = (input.text or "").strip()
        if len(text) < settings.DOCUMENT_MIN_EXTRACTED_CHARS:
            raise InsufficientJobContentError("붙여넣은 본문이 너무 짧습니다. 채용 공고 전체를 복사해 붙여주세요.")
        truncated = text[: settings.DOCUMENT_MAX_EXTRACTED_CHARS]
        return ParsedDocument(
            original_input=input,
            extracted_text=truncated,
            metadata=await self._text_metadata(input, truncated),
        )

    async def image_to_text(self, input: FileInput) -> ParsedDocument:
        """이미지 파일 파싱(미구현).

        Args:
            input: 파싱할 이미지 파일 정보.

        Raises:
            NotImplementedError: 아직 구현되지 않음.
        """
        raise NotImplementedError

    async def html_to_text(
        self,
        html: str | list[str],
        input: UrlInput,
        image_texts: list[str] | None = None,
        structured_debug: JobPostingExtractDebug | None = None,
    ) -> ParsedDocument:
        """HTML(들)에서 본문을 추출·검증해 ParsedDocument로 만든다.

        Args:
            html: 추출할 HTML 문자열 또는 여러 프레임 HTML 리스트.
            input: 원본 채용공고 URL 입력.
            image_texts: 이미지에서 추출한 보조 텍스트들.
            structured_debug: 구조화 추출 디버그 정보.

        Returns:
            추출·절단된 텍스트와 메타데이터가 담긴 ParsedDocument.

        Raises:
            InsufficientJobContentError: 추출 텍스트가 최소 길이에 못 미칠 때.
        """
        htmls = [html] if isinstance(html, str) else html
        parts: list[str] = []
        for i, h in enumerate(htmls):
            extracted = (trafilatura.extract(h) or "").strip()
            logger.info("frame[%d] html_len=%d extracted_len=%d", i, len(h), len(extracted))
            if extracted:
                parts.append(extracted)
        text = "\n\n".join(parts)
        image_texts = image_texts or []
        if image_texts:
            text = (text + "\n\n" + "\n\n".join(image_texts)).strip()

        # 이미지 텍스트까지 합친 뒤 검사 — 공고가 통이미지여도 통과
        if len(text) < settings.DOCUMENT_MIN_EXTRACTED_CHARS:
            raise InsufficientJobContentError(
                "이 URL에서 텍스트를 충분히 추출하지 못했습니다. 채용 공고 본문을 직접 복사해 붙여주세요."
            )

        truncated = text[: settings.DOCUMENT_MAX_EXTRACTED_CHARS]

        return ParsedDocument(
            original_input=input,
            extracted_text=truncated,
            metadata=self._url_metadata(input, truncated, image_texts, structured_debug),
        )
