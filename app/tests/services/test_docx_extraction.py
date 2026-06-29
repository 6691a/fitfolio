from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.schemas.documents import DocumentFormat, DocumentKind, FileInput
from app.services.document import DocumentService
from app.services.errors import DocumentFileParseError


def _make_docx(path: Path) -> None:
    document = DocxDocument()
    document.add_paragraph("백엔드 개발자 김핏폴리오")
    document.add_paragraph("경력: 5년")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "기술스택"
    table.rows[0].cells[1].text = "Python, FastAPI"
    document.save(str(path))


def test_extract_docx_text_includes_paragraphs_and_tables(tmp_path):
    docx_path = tmp_path / "resume.docx"
    _make_docx(docx_path)

    text = DocumentService._extract_docx_text(str(docx_path))

    assert "백엔드 개발자 김핏폴리오" in text
    assert "경력: 5년" in text
    assert "기술스택" in text
    assert "Python, FastAPI" in text


@pytest.mark.asyncio
async def test_docx_to_text_raises_for_non_docx(tmp_path):
    bad_path = tmp_path / "not-a.docx"
    bad_path.write_text("이건 docx가 아닙니다", encoding="utf-8")

    service = DocumentService(documents_repository=None)  # type: ignore[arg-type]
    input = FileInput(
        document_type=DocumentKind.RESUME,
        input_type=DocumentFormat.DOCX,
        file_path=str(bad_path),
        file_name="not-a.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    with pytest.raises(DocumentFileParseError):
        await service.docx_to_text(input)
